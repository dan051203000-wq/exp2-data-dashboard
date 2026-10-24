# -*- coding: utf-8 -*-
"""生成实验2数据看板实验报告 docx"""
import os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
doc = Document()

# 全局默认字体
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def set_cn(run, font="宋体", size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = RGBColor(*color)

def h(level, text):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    set_cn(r, "黑体", {1: 16, 2: 14, 3: 12.5}[level], bold=True, color=(0x1F, 0x3B, 0x63))
    return p

def para(text, bold=False, size=12, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_cn(r, "宋体", size, bold=bold)
    return p

def bullet(text, size=12):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_cn(r, "宋体", size)
    return p

def img(path, width_cm=15.5, caption=""):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption)
        set_cn(r, "楷体", 10.5, color=(0x66, 0x66, 0x66))

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htxt in enumerate(headers):
        c = t.cell(0, j)
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        set_cn(r, "黑体", 10.5, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            set_cn(r, "宋体", 10.5)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t

# ============ 封面标题 ============
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("实验2  数据看板的设计与实现"); set_cn(r, "黑体", 22, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("—— 基于 Streamlit + Pandas + Plotly 的校园一卡通消费数据看板")
set_cn(r, "楷体", 13, color=(0x55, 0x55, 0x55))

# ============ 一、实验名称 ============
h(1, "一、实验名称")
para("数据看板的设计与实现")

# ============ 二、实验目的 ============
h(1, "二、实验目的")
for s in [
    "掌握数据看板的基本概念、核心功能与应用场景，理解数据可视化、数据大屏、实时看板的区别与联系；",
    "熟悉看板设计原则与布局规范（视觉层级、色彩搭配、图表选型、信息降噪、适配布局），能够依据业务需求规划看板结构；",
    "学会依据数据特征选择适配图表（数据卡片、柱状图、折线图、饼图、仪表盘等），实现数据的直观展示与多维分析；",
    "独立完成从需求分析、版面设计、数据处理到看板落地的完整流程，具备基础的数据监控与分析能力；",
    "理解静态看板与动态实时看板的差异，掌握数据刷新、联动交互、预警监控等功能的实现原理。",
]:
    bullet(s)

# ============ 三、实验环境 ============
h(1, "三、实验环境")
table(["类别", "配置"],
      [["硬件", "个人计算机（Windows 11）、网络环境"],
       ["编辑器", "VS Code"],
       ["开发语言", "Python 3.11"],
       ["核心依赖", "Streamlit 1.66 / Pandas 3.1 / Plotly 7.1"],
       ["辅助工具", "Chrome DevTools Protocol（自动化截图）、Matplotlib（版面草图）、python-docx（报告生成）"]],
      widths=[3, 13])

# ============ 四、需求分析与版面设计 ============
h(1, "四、需求分析与版面设计（基础任务1）")
h(2, "4.1 业务场景与核心需求")
para("本实验选择「校园一卡通消费数据」作为看板业务场景，面向学校后勤与学工部门，统计周期为 2025 秋季学期（2025-09-01 ~ 2025-12-31）。看板需要回答三类业务问题：")
bullet("钱花在哪：全校消费总额、笔数、活跃人数等总量指标如何？（数据概览区）")
bullet("结构如何：各学院、各消费类型、各场景、各时段的消费分布与对比？（维度统计区）")
bullet("趋势与异常：每日消费走势如何？消费高峰集中在何时？是否存在异常波动需要预警？（趋势分析区）")
h(2, "4.2 四大模块布局规划")
table(["模块", "内容", "对应图表"],
      [["数据概览区", "核心指标卡片：消费总额 / 笔数 / 活跃人数 / 人均 / 日均 / 环比增长率", "KPI 数据卡片 ×6"],
       ["维度统计区", "学院对比、类型结构、场景排行、时段分布、支付方式", "柱状图、环形图、条形排行榜、双轴组合图"],
       ["趋势分析区", "每日消费趋势、星期×小时热力、学院雷达、人均目标", "折线图、热力图、雷达图、仪表盘"],
       ["明细数据区", "清洗后消费流水明细，支持筛选与导出", "交互式数据表格"]],
      widths=[3, 8, 5])
h(2, "4.3 配色方案")
table(["用途", "颜色", "说明"],
      [["主色", "#2B6CFF 科技蓝", "标题强调、KPI 高亮、链接色"],
       ["辅助色", "#20C997 青绿 / #FFD43B 琥珀 / #FF4D6D 警示红", "正向指标 / 中性图表 / 预警提示"],
       ["背景色", "#0D1424 深空黑（暗色）/ #F7F9FC 云白（亮色）", "暗色科技风为主，亮色简约风可一键切换"],
       ["正文色", "#EAF2FF / #8FA8CC", "主文字 / 次要说明，保证对比度"]],
      widths=[3, 6.5, 6.5])
h(2, "4.4 版面草图")
img(os.path.join(BASE, "docs", "layout_sketch.png"), 16, "图1  看板版面草图（左侧控制台 + 右侧三大图表区，底部为明细数据区）")
para("布局采用 Streamlit 经典「侧边栏 + 主区」结构：侧边栏承载全部筛选与配置（时间、维度、阈值、主题、实时开关），主区按「概览 → 维度 → 趋势 → 明细」自上而下组织，符合用户从总览到细节的阅读动线；整页采用栅格化双列/三列卡片布局，保证信息密度与留白的平衡。")

# ============ 五、数据准备与预处理 ============
h(1, "五、实验数据准备与预处理（基础任务2）")
h(2, "5.1 模拟数据生成")
para("由于真实一卡通数据不可获取，编写 generate_data.py 按真实作息规律生成模拟数据：4 个校区 × 10 个学院 × 420 名学生，覆盖早餐（7-9 点）、午餐（11-13 点）、晚餐（17-19 点）三个消费高峰，周末消费量降至 70%，并叠加噪声扰动。共生成 6040 条原始流水，同时人为注入三类脏数据用于演练预处理：")
bullet("重复记录 54 条（完全重复流水）")
bullet("空值 61 处（消费场景、支付方式字段随机缺失）")
bullet("异常值 20 条（单笔 -999 或 99999 元、学期外日期）")
h(2, "5.2 预处理流程与结果")
para("preprocess.py 按「去重 → 空值处理 → 异常值剔除 → 维度整理」四步清洗：重复流水按全字段去重；文本空值按众数填补；金额异常值采用业务规则（0 < 金额 ≤ 500 元）结合 IQR 法双重过滤；随后派生日期、小时、星期、时段、单笔等级等统计维度。清洗结果如下：")
table(["阶段", "记录数", "说明"],
      [["原始数据", "6040 条", "含 54 条重复、61 处空值、20 条异常"],
       ["去重后", "5986 条", "剔除完全重复流水"],
       ["空值填补后", "5986 条", "文本字段众数填补，无记录删除"],
       ["异常剔除后（最终）", "5905 条", "剔除金额/日期异常，保留率 97.76%"]],
      widths=[4.5, 3.5, 8])
para("清洗后的数据保存为 data/clean_campus_card_data.csv（UTF-8-SIG 编码），作为看板的唯一数据源。")

# ============ 六、看板实现 ============
h(1, "六、看板实现（基础任务3/4 + 进阶 + 扩展）")
h(2, "6.1 技术选型")
para("采用 Streamlit 作为看板框架（纯 Python、组件生态完善、自带响应式布局），Pandas 负责数据计算与筛选联动，Plotly Graph Objects 手写图表工厂实现 10 类图表组件并统一注入主题色，CSS 自定义设计系统控制卡片、字号、间距与动画。")
h(2, "6.2 任务完成清单")
table(["层级", "任务", "完成情况与实现要点"],
      [["基础", "任务1 需求分析与版面设计", "✔ 四大模块规划 + 暗色科技风配色 + 版面草图（图1）"],
       ["基础", "任务2 数据准备与预处理", "✔ 6040 条模拟流水，清洗至 5905 条（保留率 97.76%）"],
       ["基础", "任务3 基础可视化图表", "✔ KPI 卡片 ×6、柱状图、折线图、环形图（超过 3 种要求）"],
       ["基础", "任务4 静态看板整合", "✔ 统一 CSS 设计系统、圆角卡片、视觉层级、间距字号统一"],
       ["基础", "任务5 结果验证与总结", "✔ 8 张自动化截图 + 问题清单（第八节）"],
       ["进阶", "进阶1 动态视觉效果", "✔ KPI 数值滚动动画、组件渐入动画、图表 hover 提示"],
       ["进阶", "进阶2 图表联动与筛选", "✔ 五维筛选联动；点击柱状图直接下钻至对应学院"],
       ["进阶", "进阶3 多维度分析可视化", "✔ 排行榜、热力图、雷达图、目标仪表盘"],
       ["扩展", "扩展1 实时数据看板", "✔ 实时模式开关 + 模拟动态数据流 + 秒级自动刷新"],
       ["扩展", "扩展2 数据预警监控", "✔ 双阈值滑块，超限触发红色预警条 + KPI 变色"],
       ["扩展", "扩展3 大屏与移动端适配", "✔ 2560px 大屏布局 + 430px 移动端单列适配"],
       ["扩展", "扩展4 统计分析与报告导出", "✔ CSV/Excel 一键导出 + 自动生成 Markdown 分析结论"],
       ["扩展", "扩展5 个性化定制与创新", "✔ 暗/亮双主题切换、图表钻取、URL 参数直达演示状态"]],
      widths=[1.8, 5.5, 8.7])

# ============ 七、成果展示 ============
h(1, "七、看板成果展示（基础任务5 · 截图）")
shots = [
    ("01-overview-dark-fold.png", "图2  暗色科技风 · 首屏（预警阈值正常时显示运行正常徽标）"),
    ("02-overview-dark-full.png", "图3  暗色科技风 · 整页（概览/维度/趋势/明细四大区域完整形态）"),
    ("03-light-theme.png", "图4  亮色简约风主题（扩展5：一键主题切换）"),
    ("04-drilldown-college.png", "图5  图表联动下钻（进阶2：点击柱体下钻至该学院，KPI 与图表同步过滤）"),
    ("05-alert-warning.png", "图6  预警触发（扩展2：日均消费额与大单占比同时超过预警线，红色预警条+KPI 变色）"),
    ("06-realtime-mode.png", "图7  实时数据流模式（扩展1：开启实时刷新，右上角显示实时数据流已接入徽标）"),
    ("07-mobile.png", "图8  移动端适配（扩展3：430×932 视口，单列堆叠布局）"),
    ("08-bigscreen-2560.png", "图9  大屏适配（扩展3：2560px 超宽视口，多列栅格布局）"),
]
for f, cap in shots:
    img(os.path.join(BASE, "screenshots", f), 15.5, cap)

# ============ 八、问题与解决 ============
h(1, "八、实验过程中遇到的问题及解决方法")
table(["#", "问题", "原因分析", "解决方法"],
      [["1", "Plotly 7 报错 Invalid property 'titlefont'", "Plotly 7.0 移除了 titlefont 属性", "改为 title=dict(text=..., font=dict(...)) 新写法，并编写图表组件单元自测脚本逐个验证"],
       ["2", "KPI 动画在页面上重复渲染几十张卡片", "循环调用 st.markdown 追加内容而非替换", "改用 st.empty() 占位符原地更新实现滚动动画"],
       ["3", "chromedriver 无法启动，selenium 截图失败", "本机环境阻止 chromedriver 进程", "改用 Chrome DevTools Protocol（CDP）WebSocket 直连截图，不依赖 chromedriver"],
       ["4", "图表点击联动与滑块调节的自动化点击不生效", "scrollIntoView 未等完成就取坐标，点击位置漂移", "用 Promise 包裹滚动后延迟取坐标，并补发 mouseMoved 事件；同时为看板增加 URL 参数直达演示状态作为回退"],
       ["5", "卡片样式依赖 Streamlit 内部 DOM 结构，版本升级后失效", "data-testid 属性随版本变化", "放弃内部选择器，改为对 st.plotly_chart/st.dataframe 输出容器直接注入样式"],
       ["6", "新版本 Streamlit 的容器边框样式被全局 CSS 覆盖后错位", "对内部容器滥用 !important", "精简 CSS 规则，仅保留对顶层布局容器的控制"]],
      widths=[0.8, 4.2, 4.5, 6.5])

# ============ 九、总结 ============
h(1, "九、实验总结")
para("本实验完整走通了数据看板从需求分析、版面设计、数据生成与清洗、图表搭建、整合排版到成果验证的全流程。基础任务全部完成，并在进阶与扩展层面完成了 3 + 5 项：实现了五维筛选联动与图表点击下钻、双阈值预警监控、实时数据流模拟、大屏与移动端多终端适配、双主题切换以及报告导出闭环。", indent=True)
para("主要收获：一是理解了「图表选型服务于业务问题」——概览区用 KPI 卡片回答总量问题，维度区用对比与结构图表回答分布问题，趋势区用折线/热力回答时序问题；二是体会到预处理决定看板可信度，6040 条脏数据若不清洗，KPI 会被 99999 元的异常值直接拉爆；三是掌握了 Streamlit 的事件驱动模型与缓存机制，理解了联动交互本质上是「筛选条件变化 → 全页重算重绘」。", indent=True)
para("不足与改进方向：目前数据源仍为模拟数据，后续可接入真实数据库或校园 API；预警规则目前为静态阈值，可引入同比/环比动态基线；移动端图表字号在极小屏下仍偏小，可增加专门的移动端图表尺寸策略。", indent=True)

out = os.path.join(BASE, "docs", "实验2 数据看板的设计与实现-实验报告.docx")
doc.save(out)
print("saved:", out)
